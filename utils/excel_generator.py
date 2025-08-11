"""
Excel範本生成工具
用於產生費率設定的Excel範本供客服使用
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import os


class ExcelTemplateGenerator:
    """Excel範本生成器"""

    def __init__(self):
        self.wb = Workbook()

    def create_rate_plan_template(
        self, output_path: str = "templates/rate_plan_template.xlsx"
    ):
        """建立費率方案設定範本"""

        # 移除預設工作表
        if "Sheet" in self.wb.sheetnames:
            self.wb.remove(self.wb["Sheet"])

        # 建立各種時段範本工作表
        self._create_single_period_sheet()
        self._create_two_period_sheet()
        self._create_three_period_sheet()
        self._create_multi_period_sheet()
        self._create_progressive_rate_sheet()
        self._create_instruction_sheet()

        # 確保目錄存在
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        # 儲存檔案
        self.wb.save(output_path)
        print(f"Excel範本已生成: {output_path}")

    def _create_single_period_sheet(self):
        """建立一段式費率範本"""
        ws = self.wb.create_sheet("一段式費率")

        # 標題
        ws["A1"] = "一段式費率設定範本"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")

        # 欄位標題
        headers = [
            "方案ID",
            "方案名稱",
            "方案說明",
            "日期類型",
            "時段名稱",
            "開始時間",
            "結束時間",
            "計費單位(分)",
            "單位費用",
            "啟用上限",
            "上限金額",
            "免費時間(分)",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        # 範例資料
        example_data = [
            [
                "simple_rate",
                "簡單費率",
                "適用於一般露天停車場",
                "weekday",
                "全天",
                "00:00",
                "24:00",
                60,
                30,
                "是",
                240,
                15,
            ],
            ["", "", "", "holiday", "全天", "00:00", "24:00", 60, 40, "是", 320, 10],
        ]

        for row, data in enumerate(example_data, 4):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        # 設定欄寬
        column_widths = [15, 20, 30, 12, 12, 12, 12, 15, 12, 12, 12, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    def _create_two_period_sheet(self):
        """建立兩段式費率範本"""
        ws = self.wb.create_sheet("兩段式費率")

        # 標題
        ws["A1"] = "兩段式費率設定範本（日間/夜間）"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(
            start_color="70AD47", end_color="70AD47", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")

        # 欄位標題
        headers = [
            "方案ID",
            "方案名稱",
            "方案說明",
            "日期類型",
            "時段名稱",
            "開始時間",
            "結束時間",
            "計費單位(分)",
            "單位費用",
            "啟用上限",
            "上限金額",
            "免費時間(分)",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        # 範例資料
        example_data = [
            [
                "two_period_weekday",
                "平日兩段",
                "平日日夜分段收費",
                "weekday",
                "日間",
                "08:00",
                "22:00",
                60,
                30,
                "是",
                120,
                15,
            ],
            ["", "", "", "weekday", "夜間", "22:00", "08:00", 60, 10, "是", 60, 0],
            [
                "two_period_holiday",
                "假日兩段",
                "假日日夜分段收費",
                "holiday",
                "日間",
                "08:00",
                "22:00",
                60,
                40,
                "是",
                160,
                10,
            ],
            ["", "", "", "holiday", "夜間", "22:00", "08:00", 60, 20, "是", 80, 0],
        ]

        for row, data in enumerate(example_data, 4):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        # 設定欄寬
        column_widths = [18, 20, 30, 12, 12, 12, 12, 15, 12, 12, 12, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    def _create_three_period_sheet(self):
        """建立三段式費率範本"""
        ws = self.wb.create_sheet("三段式費率")

        # 標題
        ws["A1"] = "三段式費率設定範本（日間/下午/夜間）"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(
            start_color="FFC000", end_color="FFC000", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=14, color="000000")

        # 欄位標題
        headers = [
            "方案ID",
            "方案名稱",
            "方案說明",
            "日期類型",
            "時段名稱",
            "開始時間",
            "結束時間",
            "計費單位(分)",
            "單位費用",
            "啟用上限",
            "上限金額",
            "免費時間(分)",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        # 範例資料
        example_data = [
            [
                "three_period_mall",
                "商場三段",
                "商場三時段收費",
                "weekday",
                "日間",
                "08:00",
                "14:00",
                60,
                40,
                "是",
                120,
                10,
            ],
            ["", "", "", "weekday", "下午", "14:00", "18:00", 60, 50, "是", 100, 0],
            ["", "", "", "weekday", "夜間", "18:00", "08:00", 60, 20, "是", 80, 0],
        ]

        for row, data in enumerate(example_data, 4):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        # 設定欄寬
        column_widths = [18, 20, 30, 12, 12, 12, 12, 15, 12, 12, 12, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    def _create_multi_period_sheet(self):
        """建立多段式費率範本"""
        ws = self.wb.create_sheet("多段式費率")

        # 標題
        ws["A1"] = "多段式費率設定範本（自定義時段）"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(
            start_color="C55A5A", end_color="C55A5A", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")

        # 欄位標題
        headers = [
            "方案ID",
            "方案名稱",
            "方案說明",
            "日期類型",
            "時段名稱",
            "開始時間",
            "結束時間",
            "計費單位(分)",
            "單位費用",
            "啟用上限",
            "上限金額",
            "免費時間(分)",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="F2DCDB", end_color="F2DCDB", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        # 範例資料
        example_data = [
            [
                "multi_period_airport",
                "機場多段",
                "機場依流量分段",
                "weekday",
                "早晨",
                "06:00",
                "09:00",
                30,
                15,
                "否",
                0,
                30,
            ],
            ["", "", "", "weekday", "上午尖峰", "09:00", "12:00", 30, 25, "是", 60, 0],
            ["", "", "", "weekday", "午餐", "12:00", "14:00", 60, 40, "是", 80, 0],
            ["", "", "", "weekday", "下午尖峰", "14:00", "18:00", 30, 30, "是", 120, 0],
            ["", "", "", "weekday", "晚間", "18:00", "22:00", 60, 35, "是", 100, 0],
            ["", "", "", "weekday", "夜間", "22:00", "06:00", 60, 10, "是", 50, 0],
        ]

        for row, data in enumerate(example_data, 4):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        # 設定欄寬
        column_widths = [20, 20, 30, 12, 15, 12, 12, 15, 12, 12, 12, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    def _create_progressive_rate_sheet(self):
        """建立累進費率設定範本"""
        ws = self.wb.create_sheet("累進費率設定")

        # 標題
        ws["A1"] = "累進費率設定範本"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(
            start_color="7030A0", end_color="7030A0", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")

        # 說明
        ws["A2"] = "累進費率：依停車時間長短使用不同費率，如前1小時$30，第2小時起$40"
        ws["A2"].font = Font(size=12, color="7030A0")

        # 欄位標題
        headers = [
            "方案ID",
            "時段名稱",
            "階層",
            "起始分鐘",
            "結束分鐘",
            "計費單位(分)",
            "單位費用",
            "備註",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="E1D5ED", end_color="E1D5ED", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        # 範例資料
        example_data = [
            ["progressive_mall", "上午尖峰", "第1階", 0, 60, 30, 25, "前1小時"],
            ["", "", "第2階", 60, "無限制", 30, 35, "第2小時起"],
            ["progressive_office", "辦公時段", "第1階", 0, 120, 60, 40, "前2小時"],
            ["", "", "第2階", 120, 300, 60, 50, "第3-5小時"],
            ["", "", "第3階", 300, "無限制", 60, 60, "第6小時起"],
        ]

        for row, data in enumerate(example_data, 5):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        # 設定欄寬
        column_widths = [18, 15, 12, 12, 12, 15, 12, 20]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    def _create_instruction_sheet(self):
        """建立使用說明工作表"""
        ws = self.wb.create_sheet("使用說明")
        ws.sheet_view.show_grid_lines = False

        # 標題
        ws["A1"] = "停車費率設定系統 - 使用說明"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A1"].fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")

        # 說明內容
        instructions = [
            "",
            "一、系統支援的時段類型：",
            "  1. 一段式：全天統一費率，適用於簡單停車場",
            "  2. 兩段式：日間/夜間分段，最常見的收費方式",
            "  3. 三段式：日間/下午/夜間，適用於商場或辦公區",
            "  4. 多段式：自定義多個時段，適用於複雜收費需求",
            "",
            "二、欄位說明：",
            "  • 方案ID：英文代碼，用於系統識別（必填）",
            "  • 方案名稱：中文顯示名稱（必填）",
            "  • 日期類型：weekday=平日, holiday=假日（必填）",
            "  • 開始/結束時間：24小時制，如08:00（必填）",
            "  • 計費單位：幾分鐘為一個收費單位（必填）",
            "  • 單位費用：每個單位的收費金額（必填）",
            "  • 啟用上限：是/否，該時段是否有收費上限",
            "  • 上限金額：該時段最高收費（啟用上限時必填）",
            "  • 免費時間：前幾分鐘免費，0表示無免費時間",
            "",
            "三、特殊功能：",
            "  • 跨日時段：結束時間小於開始時間，如22:00-08:00",
            "  • 累進費率：依停車時間長短採用不同費率",
            "  • 免費緩衝：前幾分鐘免費，減少短暫停車收費",
            "  • 收費上限：避免長時間停車產生天價費用",
            "",
            "四、範例場景：",
            "  • 住宅區：夜間免費或低價，日間正常收費",
            "  • 商業區：上班時間高價，其他時段低價",
            "  • 購物中心：消費時段分級收費，深夜低價",
            "  • 轉運站：短停免費，長停累進收費",
            "",
            "五、注意事項：",
            "  • 同一方案的時段不可重疊",
            "  • 跨日時段會自動處理日期切換",
            "  • 免費時間從停車開始計算",
            "  • 上限金額僅適用於該時段，不是全天上限",
        ]

        for i, instruction in enumerate(instructions, 2):
            ws[f"A{i}"] = instruction
            if instruction.startswith(("一、", "二、", "三、", "四、", "五、")):
                ws[f"A{i}"].font = Font(bold=True, size=12)

        # 設定欄寬
        ws.column_dimensions["A"].width = 80


def main():
    """生成Excel範本"""
    generator = ExcelTemplateGenerator()
    generator.create_rate_plan_template()


if __name__ == "__main__":
    main()
